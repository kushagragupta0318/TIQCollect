# ─── CHANGELOG (prototype → product) ───
# New file, 2026-09-24. E09 — app/reports/: payload, formatting, the three
# renderers and the export seam.
#
# No database, no MinIO, no network. Every renderer test renders the sample
# payload and READS THE FILE BACK with the library that wrote the format
# (python-pptx, openpyxl) or, for the PDF, from its inflated streams — the
# assertion is on what a reader of the file would get, not on what the code
# meant to write. Organisation names are read from app/reports/sample.py and
# never spelled out here: the demo bank's name is under review.
import io
import re
import zlib

import pytest

from app.reports import formatting as f
from app.reports import sample, theme
from app.reports.payload import Chart, ChartKind, Column, ReportPayload, Series, Table, Unit

RUPEE_CHAR = chr(0x20B9)
NOT_IN_FONT = chr(0xE000)          # private use: Noto Sans has no glyph for it


@pytest.fixture(scope="module")
def payload() -> ReportPayload:
    return sample.sample_payload()


# ── formatting ───────────────────────────────────────────────────────────────
@pytest.mark.parametrize("n,expected", [
    (0, "0"), (999, "999"), (1000, "1,000"), (99_999, "99,999"), (1_00_000, "1,00,000"),
    (12_34_567, "12,34,567"), (1_23_45_678, "1,23,45,678"), (-45_00_000, "-45,00,000"),
])
def test_indian_grouping(n, expected):
    assert f.indian_grouping(n) == expected


def test_rupee_amounts_full_and_compact():
    assert f.inr(1_50_000) == RUPEE_CHAR + "1,50,000"
    assert f.inr(1_50_000, compact=True) == RUPEE_CHAR + "1.50 L"
    assert f.inr(3_08_40_000, compact=True) == RUPEE_CHAR + "3.08 Cr"
    assert f.inr(-25_00_000, compact=True) == "-" + RUPEE_CHAR + "25.00 L"
    assert f.inr(9_999, compact=True) == RUPEE_CHAR + "9,999"


def test_a_missing_value_is_a_dash_never_zero():
    for unit in Unit:
        if unit is not Unit.TEXT:
            assert f.fmt(None, unit) == f.MISSING
    assert f.fmt(float("nan"), Unit.INR) == f.MISSING


def test_percentages_are_fractions_and_change_in_points():
    assert f.fmt(0.2468, Unit.PCT) == "24.7%"
    assert f.delta(0.247, 0.192, Unit.PCT) == "+5.5 pp"
    assert f.delta(110.0, 100.0, Unit.INR) == "+10.0%"
    assert f.delta(5.0, 0.0, Unit.COUNT) is None       # no relative change from zero
    assert f.delta(None, 1.0, Unit.INR) is None


def test_one_scale_per_rupee_axis():
    assert f.axis_scale(6.66e7, Unit.INR) == (1e7, "Cr")
    assert f.axis_scale(4.2e6, Unit.INR) == (1e5, "L")
    assert f.axis_scale(9_000, Unit.INR) == (1.0, "")
    # the defect this fixed: half a crore printed as "₹50.00 L" on a crore axis
    assert f.axis_tick(5e6, Unit.INR, 1e7, "Cr") == RUPEE_CHAR + "0.5 Cr"
    assert f.axis_tick(0.175, Unit.PCT, 1.0, "") == "17.5%"


def test_period_label(payload):
    assert f.period_label(payload) == "1 Sep 2026 – 30 Sep 2026"


# ── payload validation ───────────────────────────────────────────────────────
def test_a_table_row_cannot_carry_an_undeclared_column():
    with pytest.raises(ValueError, match="no column declares"):
        Table(table_id="t", title="t", columns=[Column(key="a", label="A")], rows=[{"a": 1, "b": 2}])


def test_a_series_must_fit_its_categories():
    with pytest.raises(ValueError, match="2 values for 3 categories"):
        Chart(chart_id="c", title="c", kind=ChartKind.BAR, unit=Unit.COUNT,
              categories=["x", "y", "z"], series=[Series(name="s", values=[1, 2])])


def test_no_chart_may_need_a_generated_colour():
    too_many = [Series(name=f"s{i}", values=[1.0]) for i in range(len(theme.SERIES) + 1)]
    with pytest.raises(ValueError, match="at most"):
        Chart(chart_id="c", title="c", kind=ChartKind.BAR, unit=Unit.COUNT, categories=["x"], series=too_many)
    # a reference line takes no palette slot
    Chart(chart_id="c", title="c", kind=ChartKind.LINE, unit=Unit.COUNT, categories=["x"],
          series=[*too_many[:len(theme.SERIES)], Series(name="target", values=[1.0], reference=True)])


def test_ids_are_unique_across_the_payload(payload):
    s = payload.sections[0]
    with pytest.raises(ValueError, match="duplicate id"):
        ReportPayload(**{**payload.model_dump(), "sections": [s.model_dump(), s.model_dump()]})


# ── theme ────────────────────────────────────────────────────────────────────
def test_series_colours_follow_position_and_references_are_muted(payload):
    trend = next(c for s in payload.sections for c in s.charts if c.kind is ChartKind.LINE)
    assert theme.series_colors(trend) == [theme.SERIES[0], theme.REFERENCE]


def test_ordered_series_use_one_hue_light_to_dark(payload):
    dpd = next(c for s in payload.sections for c in s.charts if c.scale == "sequential")
    cols = theme.series_colors(dpd)
    idx = [theme.SEQUENTIAL.index(c) for c in cols]
    assert idx == sorted(idx) and len(set(idx)) == len(idx)


# ── the rupee glyph actually renders from the bundled font ───────────────────
def test_the_bundled_fonts_ship_with_their_licence():
    assert theme.FONT_REGULAR.is_file() and theme.FONT_BOLD.is_file()
    assert "SIL Open Font License" in (theme.FONT_DIR / "OFL.txt").read_text(encoding="utf-8")


def test_the_bundled_font_draws_a_rupee_not_a_missing_glyph_box():
    """Rasterise the rupee sign and a code point the font does not have. A font
    without the rupee falls back to .notdef — the same box both times."""
    import numpy as np
    from fontTools.ttLib import TTFont
    from matplotlib.ft2font import FT2Font

    font = FT2Font(str(theme.FONT_REGULAR))
    font.set_size(48, 72)

    def raster(ch: str):
        font.set_text(ch, 0.0)
        font.draw_glyphs_to_bitmap()
        return np.asarray(font.get_image()).copy()

    rupee, missing = raster(RUPEE_CHAR), raster(NOT_IN_FONT)
    assert rupee.sum() > 0, "the rupee glyph drew nothing"
    assert rupee.shape != missing.shape or not np.array_equal(rupee, missing), (
        "the rupee drew the same box as a code point the font does not have")
    cmap = TTFont(str(theme.FONT_REGULAR)).getBestCmap()
    assert 0x20B9 in cmap and 0xE000 not in cmap


# ── PDF ──────────────────────────────────────────────────────────────────────
# ReportLab ends a stream's data directly with "endstream" (no newline), and
# encodes content streams ASCII85 + Flate, font streams Flate only.
_STREAM = re.compile(rb"stream\r?\n(.*?)endstream", re.S)


def _a85(data: bytes) -> bytes | None:
    import base64
    s = data.strip()
    s = s[2:] if s.startswith(b"<~") else s
    s = s[:-2] if s.endswith(b"~>") else s
    try:
        return base64.a85decode(s)
    except ValueError:
        return None


def _inflate(data: bytes) -> bytes | None:
    for candidate in (data, _a85(data)):
        if candidate is None:
            continue
        try:
            return zlib.decompressobj().decompress(candidate)      # ignores trailing bytes
        except zlib.error:
            continue
    return None


@pytest.fixture(scope="module")
def pdf_bytes(payload) -> bytes:
    """The PDF as shipped, with every stream decoded and appended, so the
    assertions read fonts and CMaps whatever the compression setting."""
    from app.reports.render_pdf import render_pdf
    raw = render_pdf(payload)
    parts = [raw, *filter(None, (_inflate(m.group(1)) for m in _STREAM.finditer(raw)))]
    return b"\n".join(parts)


def test_pdf_is_a_pdf_with_several_pages(pdf_bytes):
    assert pdf_bytes.startswith(b"%PDF-")
    pages = len(re.findall(rb"/Type /Page[^s]", pdf_bytes))
    assert pages >= 3, pages   # cover + sections, and the 30-row table crosses a page


def test_pdf_embeds_noto_sans_and_uses_its_rupee_glyph(pdf_bytes):
    assert b"/FontFile2" in pdf_bytes                     # a TrueType font is embedded
    assert re.search(rb"/BaseFont /(?:[A-Z]{6}\+)?NotoSans", pdf_bytes)
    # Every glyph drawn is listed in the embedded font's ToUnicode map. The
    # rupee appears there only if a rupee sign was actually set in Noto Sans.
    assert re.search(rb"<20[bB]9>", pdf_bytes), "no rupee glyph was drawn with the embedded font"


def test_pdf_carries_the_organisation_in_its_metadata(pdf_bytes, payload):
    assert payload.organisation.encode() in pdf_bytes


def test_pdf_renders_every_chart(payload, pdf_bytes):
    n_charts = sum(len(s.charts) for s in payload.sections)
    assert len(re.findall(rb"/Subtype /Image", pdf_bytes)) >= n_charts


# ── PPTX ─────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def deck(payload):
    from pptx import Presentation
    from app.reports.render_pptx import render_pptx
    return Presentation(io.BytesIO(render_pptx(payload)))


def _slide_text(slide) -> str:
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_pptx_title_slide_names_the_organisation_and_warns(deck, payload):
    first = _slide_text(deck.slides[0])
    assert payload.organisation.upper() in first
    assert "synthetic" in first.lower()


def test_pptx_charts_are_native_and_carry_the_payload_values(deck, payload):
    from pptx.enum.chart import XL_CHART_TYPE
    # python-pptx INFERS chart_type from the XML: a line chart whose markers
    # are set per series (as ours are — none on the reference line) reads LINE.
    expected = {ChartKind.BAR: {XL_CHART_TYPE.COLUMN_CLUSTERED},
                ChartKind.STACKED_BAR: {XL_CHART_TYPE.COLUMN_STACKED},
                ChartKind.LINE: {XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS}}
    native = [sh.chart for s in deck.slides for sh in s.shapes if sh.has_chart]
    specs = [c for s in payload.sections for c in s.charts]
    assert len(native) == len(specs)
    for chart, spec in zip(native, specs):
        assert chart.chart_type in expected[spec.kind]
        assert list(chart.plots[0].categories) == spec.categories
        peak = max(abs(v) for s in spec.series for v in s.values if v is not None)
        if spec.kind is ChartKind.STACKED_BAR:
            peak = max(sum(v or 0 for v in col) for col in zip(*(s.values for s in spec.series)))
        divisor, _ = f.axis_scale(peak, spec.unit)
        for got, want in zip(chart.plots[0].series, spec.series):
            assert got.name == want.name
            for g, w in zip(got.values, want.values):
                assert (g is None and w is None) or g == pytest.approx(w / divisor)


def test_pptx_rupee_charts_say_their_scale_in_the_title(deck):
    titles = [_slide_text(s) for s in deck.slides if any(sh.has_chart for sh in s.shapes)]
    assert any(RUPEE_CHAR + " crore" in t for t in titles)


def test_pptx_splits_a_long_table_and_loses_no_row(deck, payload):
    from app.reports.render_pptx import TABLE_ROWS_PER_SLIDE
    for t in (t for s in payload.sections for t in s.tables):
        tables = [sh.table for s in deck.slides for sh in s.shapes
                  if sh.has_table and _slide_text(s).startswith(t.title)]
        assert sum(len(x.rows) - 1 for x in tables) == len(t.rows)
        assert len(tables) == -(-len(t.rows) // TABLE_ROWS_PER_SLIDE)


def test_pptx_labels_ai_commentary_and_only_that(deck, payload):
    ai_titles = [s.title for s in payload.sections if s.narrative and s.narrative.ai_generated]
    labelled = [_slide_text(s) for s in deck.slides if "AI-WRITTEN" in _slide_text(s)]
    assert len(labelled) == len(ai_titles) >= 1
    assert all(text.startswith(title) for text, title in zip(labelled, ai_titles))


# ── XLSX ─────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def book(payload):
    from openpyxl import load_workbook
    from app.reports.render_xlsx import render_xlsx
    return load_workbook(io.BytesIO(render_xlsx(payload)))


def test_xlsx_has_a_sheet_per_table_and_per_chart(book, payload):
    n_tables = sum(len(s.tables) for s in payload.sections)
    n_charts = sum(len(s.charts) for s in payload.sections)
    assert book.sheetnames[0] == "Report"
    assert len(book.sheetnames) == 1 + n_tables + n_charts
    assert all(len(n) <= 31 for n in book.sheetnames)


def test_xlsx_numbers_are_numbers_in_their_unit_format(book, payload):
    t = next(t for s in payload.sections for t in s.tables if t.table_id == "agency_scorecard")
    ws = book[t.title]
    assert [c.value for c in ws[1]] == [c.label for c in t.columns]
    for r, row in enumerate(t.rows, start=2):
        for ci, col in enumerate(t.columns, start=1):
            cell = ws.cell(row=r, column=ci)
            want = row[col.key]
            # raw value, not formatted text — to Excel's own 15 significant digits
            assert cell.value == (pytest.approx(want, rel=1e-14) if isinstance(want, float) else want)
            if row[col.key] is not None and col.unit is not Unit.TEXT:
                assert cell.number_format == f.XLSX_FORMATS[col.unit]


def test_xlsx_report_sheet_lists_every_kpi_with_its_catalog_id(book, payload):
    ws = book["Report"]
    ids = {c.value for c in ws["C"]}
    for k in (k for s in payload.sections for k in s.kpis):
        assert k.kpi_id in ids
    rows = [r for r in ws.iter_rows(values_only=True) if r[2] == "cost_to_collect"]
    assert rows and rows[0][3] is None                   # not measurable → empty, never 0


def test_xlsx_sheet_names_stay_legal_and_unique():
    from app.reports.render_xlsx import sheet_name
    used: set[str] = set()
    long = "Recovery by agency and zone for the period: actual/expected"
    a, b = sheet_name(long, used), sheet_name(long, used)
    assert a != b and len(a) <= 31 and len(b) <= 31
    assert not re.search(r"[\[\]:*?/\\]", a + b)


# ── names come from one place ────────────────────────────────────────────────
def test_no_renderer_or_theme_spells_out_a_demo_organisation():
    """The demo bank's name is under review. It must change in sample.py and
    nowhere else — so no other module in app/reports may contain it."""
    from pathlib import Path
    names = [sample.DEMO_BANK, *(n for n, _ in sample.DEMO_AGENCIES)]
    root = Path(sample.__file__).parent
    offenders = [(p.name, n) for p in root.glob("*.py") if p.name != "sample.py"
                 for n in names if n in p.read_text(encoding="utf-8")]
    assert offenders == []


# ── the export seam ──────────────────────────────────────────────────────────
@pytest.fixture
def seams(monkeypatch):
    from app.reports import service
    calls: list[tuple] = []
    monkeypatch.setattr(service.storage, "upload_bytes",
                        lambda key, data, ct: calls.append(("upload", key, len(data), ct)))
    monkeypatch.setattr(service.storage, "presigned_download_url",
                        lambda key, expires_minutes: calls.append(("presign", key)) or f"https://files.example/{key}")
    monkeypatch.setattr(service, "write_audit", lambda db, **kw: calls.append(("audit", kw)) or True)
    return service, calls


@pytest.mark.parametrize("fmt", ["pdf", "pptx", "xlsx"])
def test_export_uploads_then_signs_then_audits(seams, payload, fmt):
    service, calls = seams
    out = service.export(None, payload, fmt, user_id="user-1", endpoint="/bank/reports/export")
    assert [c[0] for c in calls] == ["upload", "presign", "audit"]
    _, key, size, ct = calls[0]
    assert key == out.key and key.startswith(f"reports/{payload.report_id}/") and key.endswith(f".{fmt}")
    assert size == out.size_bytes > 0 and ct == service.FORMATS[fmt][1]
    audit = calls[2][1]
    assert audit["action"].value == "DATA_EXPORT" and audit["user_id"] == "user-1"
    assert audit["details"]["format"] == fmt and audit["details"]["bytes"] == size
    assert audit["details"]["sha256"] == out.sha256
    # the row records the export's shape, never its content
    assert all(not isinstance(v, (bytes, bytearray)) and (not isinstance(v, str) or len(v) < 200)
               for v in audit["details"].values())


def test_a_failed_upload_writes_no_audit_row(seams, payload, monkeypatch):
    service, calls = seams

    def boom(*_a, **_k):
        raise ConnectionError("minio unreachable")
    monkeypatch.setattr(service.storage, "upload_bytes", boom)
    with pytest.raises(ConnectionError):
        service.export(None, payload, "xlsx", user_id="user-1")
    assert calls == []


def test_a_failed_link_writes_no_audit_row(seams, payload, monkeypatch):
    service, calls = seams

    def boom(*_a, **_k):
        raise RuntimeError("cannot sign")
    monkeypatch.setattr(service.storage, "presigned_download_url", boom)
    with pytest.raises(RuntimeError):
        service.export(None, payload, "xlsx", user_id="user-1")
    assert [c[0] for c in calls] == ["upload"]


def test_an_unknown_format_is_refused_before_anything_happens(seams, payload):
    service, calls = seams
    with pytest.raises(ValueError, match="unknown report format"):
        service.export(None, payload, "docx", user_id="user-1")
    assert calls == []
