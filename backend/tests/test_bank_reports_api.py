"""POST /api/v1/bank/reports/generate and GET /api/v1/bank/reports/types
(E10, the Board Reports page's API).

The report ENGINE (payload -> PDF/PPTX/XLSX bytes) is covered end to end by
tests/test_report_engine.py; this file is the thin route layer: the two
templates actually build a valid file from a real (if view-less, SQLite) bank
database, the capability gate, and tenant scoping — an agency from another
bank reads exactly like one that does not exist, same as every other bank
route's own `_own()` check.

MinIO is faked the same way tests/test_report_engine.py fakes it
(`service.storage.*`); nothing here needs a running MinIO or network.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import insert

from app.core.database import get_db
from app.core.dependencies import _get_token_payload, get_current_user, get_tenant_analytics_db
from app.main import app
from app.models.audit_log import AuditAction, AuditLog
from app.models.tenancy import Agency, Bank
from app.models.user import User, UserRole
from app.reports import service as report_service
from tests._db import DEFAULT_TENANT, TEST_AGENCY_ID, TEST_BANK_ID, create_schema, make_engine, \
    make_session_factory, test_id

BASE = "/api/v1/bank/reports"


@pytest.fixture()
def db():
    engine = make_engine()
    create_schema(engine)
    s = make_session_factory(engine)()
    try:
        yield s
    finally:
        s.close()


@pytest.fixture()
def uploads(monkeypatch):
    """Fakes MinIO (same seam test_report_engine.py fakes) and records every
    upload, so a test can read the bytes straight back and check they are a
    real PDF/PPTX/XLSX — never that storage or the renderer were even called."""
    calls: list[tuple[str, bytes, str]] = []
    monkeypatch.setattr(report_service.storage, "upload_bytes",
                        lambda key, data, ct: calls.append((key, data, ct)))
    monkeypatch.setattr(report_service.storage, "presigned_download_url",
                        lambda key, expires_minutes: f"https://files.example/{key}")
    return calls


def _user(db, role: UserRole, **tenant) -> User:
    u = User(email=f"{role.value.lower()}-{test_id(role.value)[:8]}@girivanfinance.test",
             phone=f"98765{len(role.value):05d}", full_name="Test Person", hashed_password="x",
             role=role, is_active=True, **tenant)
    db.add(u)
    db.commit()
    return u


def _client(db, user: User) -> TestClient:
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[_get_token_payload] = lambda: {"sid": None}
    app.dependency_overrides[get_tenant_analytics_db] = lambda: db
    return TestClient(app)


@pytest.fixture()
def cleanup():
    yield
    app.dependency_overrides.clear()


def _other_bank_agency(db) -> str:
    """A second bank + agency, inserted the same way create_schema seeds the
    default one (raw Core insert, bypassing the tenant listener) — the
    cross-tenant fixture every scoping test in this codebase builds this way."""
    other_bank = test_id("bank:other-lender")
    other_agency = test_id("agency:other-lender-field")
    db.execute(insert(Bank.__table__), [{
        "id": other_bank, "code": "OTH", "legal_name": "Other Lender Finance Ltd.",
        "display_name": "Other Lender Finance", "timezone": "Asia/Kolkata", "brand": {},
        "status": "ACTIVE", "is_demo": True,
    }])
    db.execute(insert(Agency.__table__), [{
        "id": other_agency, "bank_id": other_bank, "code": "AGENCY-OTH-001",
        "legal_name": "Other Lender Field Services Pvt. Ltd.", "trade_name": "Other Lender Field Services",
        "status": "ACTIVE", "contacts": [], "is_demo": True,
    }])
    db.commit()
    return other_agency


def test_list_types_returns_board_and_agency_review(db, cleanup):
    user = _user(db, UserRole.BANK_ANALYST, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).get(f"{BASE}/types")
    assert r.status_code == 200, r.text
    templates = {t["template"]: t for t in r.json()}
    assert set(templates) == {"board", "agency_review"}
    assert templates["board"]["requires_agency"] is False
    assert templates["agency_review"]["requires_agency"] is True


def test_board_pack_generates_a_real_pdf_and_writes_a_scoped_audit_row(db, cleanup, uploads):
    user = _user(db, UserRole.BANK_ADMIN, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate", json={"template": "board", "format": "pdf"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["report_id"].startswith(f"board-{TEST_BANK_ID}-")
    assert body["format"] == "pdf" and body["size_bytes"] > 0
    assert body["url"] == f"https://files.example/{uploads[0][0]}"

    key, data, content_type = uploads[0]
    assert key.startswith(f"reports/{body['report_id']}/") and content_type == "application/pdf"
    assert data[:5] == b"%PDF-"                 # a real PDF, not a stub

    rows = db.query(AuditLog).filter(AuditLog.action == AuditAction.DATA_EXPORT).all()
    assert len(rows) == 1 and rows[0].bank_id == TEST_BANK_ID and rows[0].user_id == user.id
    assert rows[0].details["template"] == "board" and rows[0].details["format"] == "pdf"
    assert rows[0].details["sha256"] == body["sha256"]


def test_board_pack_xlsx_is_a_workbook_a_spreadsheet_app_can_open(db, cleanup, uploads):
    user = _user(db, UserRole.BANK_ADMIN, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate", json={"template": "board", "format": "xlsx"})
    assert r.status_code == 200, r.text
    _, data, content_type = uploads[0]
    assert content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames    # at least one real sheet, not an empty/corrupt file


def test_agency_review_pptx_is_a_deck_a_slide_app_can_open(db, cleanup, uploads):
    user = _user(db, UserRole.BANK_ANALYST, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate",
                               json={"template": "agency_review", "format": "pptx", "agency_id": TEST_AGENCY_ID})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["report_id"].startswith(f"agency_review-{TEST_BANK_ID}-{TEST_AGENCY_ID}-")
    _, data, content_type = uploads[0]
    assert content_type.endswith("presentationml.presentation")
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) > 0


def test_agency_review_without_an_agency_id_is_refused(db, cleanup):
    user = _user(db, UserRole.BANK_ADMIN, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate", json={"template": "agency_review", "format": "pdf"})
    assert r.status_code == 422


def test_a_malformed_agency_id_is_refused_not_a_500(db, cleanup):
    """agency_id used to be a bare str: db.get(Agency, "garbage") hits a uuid
    column and DataErrors on Postgres (the same class the actor_id fix
    728c66c closed on the audit routes). SQLite's test harness would not
    catch this on its own -- the id is now validated before any query runs,
    so this test holds on either backend."""
    user = _user(db, UserRole.BANK_ADMIN, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate",
                               json={"template": "agency_review", "format": "pdf", "agency_id": "not-a-uuid"})
    assert r.status_code == 422


def test_an_agency_from_another_bank_reads_as_not_found(db, cleanup):
    other_agency = _other_bank_agency(db)
    user = _user(db, UserRole.BANK_ADMIN, bank_id=DEFAULT_TENANT["bank_id"])
    r = _client(db, user).post(f"{BASE}/generate",
                               json={"template": "agency_review", "format": "pdf", "agency_id": other_agency})
    assert r.status_code == 404


@pytest.mark.parametrize("role", [UserRole.AGENCY_MANAGER, UserRole.AGENCY_ADMIN, UserRole.FIELD_AGENT])
def test_agency_and_field_roles_are_refused(db, cleanup, role):
    user = _user(db, role, **DEFAULT_TENANT)
    r = _client(db, user).post(f"{BASE}/generate", json={"template": "board", "format": "pdf"})
    assert r.status_code == 403
